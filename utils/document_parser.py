"""
公共文档解析模块 —— 统一 .txt/.md/.docx/.pdf/.xlsx 解析与文本分块。

供 embedding_index.py 与 rag_engine.py 共用，消除重复代码。
"""

import hashlib
import os
import re

CHUNK_MIN_CHARS = 30
CHUNK_MAX_CHARS = 1500
MAX_FILE_SIZE_MB = 5
MAX_FILES = 200

DOC_PATTERNS = {
    '.txt': True,
    '.md': True,
    '.docx': True,
    '.pdf': True,
    '.xlsx': True,
}


def _read_txt(path: str) -> str:
    for enc in ['utf-8', 'gbk', 'gb2312', 'utf-16']:
        try:
            with open(path, 'r', encoding=enc) as f:
                return f.read()
        except (UnicodeDecodeError, UnicodeError):
            continue
    return ''


def _read_docx(path: str) -> str:
    try:
        from docx import Document
        doc = Document(path)
        return '\n'.join(p.text for p in doc.paragraphs if p.text.strip())
    except ImportError:
        return ''
    except Exception:
        return ''


def _read_pdf(path: str) -> str:
    try:
        import pdfplumber
        with pdfplumber.open(path) as pdf:
            texts = [p.extract_text() for p in pdf.pages]
            return '\n'.join(t for t in texts if t)
    except (ImportError, Exception):
        pass
    try:
        from PyPDF2 import PdfReader
        reader = PdfReader(path)
        return '\n'.join(p.extract_text() or '' for p in reader.pages)
    except (ImportError, Exception):
        pass
    return ''


def _read_xlsx(path: str) -> str:
    try:
        import openpyxl
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
        lines = []
        for sheet in wb.worksheets:
            for row in sheet.iter_rows(values_only=True):
                cells = [str(c) for c in row if c is not None]
                if cells:
                    lines.append('  '.join(cells))
        return '\n'.join(lines)
    except (ImportError, Exception):
        pass
    try:
        import pandas as pd
        sheets = pd.read_excel(path, sheet_name=None)
        lines = []
        for name, df in sheets.items():
            lines.append(f'[工作表: {name}]')
            for _, row in df.iterrows():
                cells = [str(c) for c in row if pd.notna(c)]
                if cells:
                    lines.append('  '.join(cells))
        return '\n'.join(lines)
    except (ImportError, Exception):
        pass
    return ''


def read_document(path: str) -> str:
    ext = os.path.splitext(path)[1].lower()
    if ext in ('.txt', '.md'):
        return _read_txt(path)
    elif ext == '.docx':
        return _read_docx(path)
    elif ext == '.pdf':
        return _read_pdf(path)
    elif ext == '.xlsx':
        return _read_xlsx(path)
    return ''


def is_heading_line(stripped: str) -> bool:
    return bool(
        re.match(r'^(第[一二三四五六七八九十\d]+[章节条]|一、|二、|三、|四、|五、|'
                 r'六、|七、|八、|九、|十、|（[一二三四五六七八九十\d]）)', stripped)
    )


def chunk_text(text: str, source_file: str = '') -> list[dict]:
    chunks = []
    lines = text.split('\n')
    current_buf = []
    current_heading = ''
    current_len = 0

    def _flush():
        nonlocal current_len
        if current_len >= CHUNK_MIN_CHARS:
            chunks.append({
                'text': '\n'.join(current_buf).strip(),
                'source': os.path.basename(source_file) if source_file else '',
                'heading': current_heading,
            })
        current_buf.clear()
        current_len = 0

    for line in lines:
        stripped = line.strip()
        if not stripped:
            continue
        if is_heading_line(stripped) and len(stripped) < 80:
            _flush()
            current_heading = stripped
            continue
        current_buf.append(stripped)
        current_len += len(stripped)
        if current_len >= CHUNK_MAX_CHARS:
            _flush()
    _flush()
    return chunks


def file_hash(path: str) -> str:
    h = hashlib.md5()
    with open(path, 'rb') as f:
        for chunk in iter(lambda: f.read(8192), b''):
            h.update(chunk)
    return h.hexdigest()
